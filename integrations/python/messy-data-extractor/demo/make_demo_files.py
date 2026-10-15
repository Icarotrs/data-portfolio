"""Generates four fictitious supplier files, each with typical problems:

  nordhaus_pricelist.csv   semicolons, European decimals ("1.234,50 €"), cp1252
                           encoding, day-first dates ("03.04.2026"), padded text
  pacific_stock.xlsx       header in two rows with merged cells, "$1,299.00" as
                           text, a mix of real Excel dates and "Apr 3" text
  lakeside_catalog.pdf     a text-based PDF table over two pages (header repeated),
                           "Size / Color" in one column
  maple_supplies.csv       mixed date formats in one column, name split across
                           two columns, lowercase / UPPERCASE names, a duplicated
                           SKU (older and newer row), thousands separators

Deterministic: same seed → same files.  python demo/make_demo_files.py
"""

from __future__ import annotations

import csv
import random
import sys
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "demo_data" / "incoming"

PRODUCTS = [
    "stainless steel bottle 750ml", "bamboo cutting board", "café mug 350ml", "linen napkins set",
    "glass storage jar 1l", "cast iron skillet", "ceramic plate set", "wooden spoon set",
    "silicone baking mat", "cotton apron", "insulated lunch bag", "copper measuring cups",
]


def _rng(seed: int) -> random.Random:
    return random.Random(seed)


def nordhaus(rng: random.Random, base: date) -> None:
    rows = [["Art.-Nr.", "Bezeichnung", "Preis", "Bestand", "Datum"]]
    for i, p in enumerate(PRODUCTS[:8]):
        price = rng.uniform(4, 1400)
        whole, cents = f"{price:.2f}".split(".")
        whole = f"{int(whole):,}".replace(",", ".")
        d = base - timedelta(days=rng.randint(0, 20))
        rows.append([
            f" nh-{100 + i} ",
            ("  " + p.upper() + " ") if i % 3 == 0 else p,
            f"{whole},{cents} €",
            str(rng.randint(0, 400)),
            d.strftime("%d.%m.%Y"),
        ])
    text = "\r\n".join(";".join(r) for r in rows) + "\r\n"
    (OUT / "nordhaus_pricelist.csv").write_bytes(text.encode("cp1252"))


def pacific(rng: random.Random, base: date) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Item", None, "Pricing", None, "Inventory", None])
    ws.append(["Code", "Description", "Unit price", "Currency", "On hand", "Updated"])
    for a, b in (("A1", "B1"), ("C1", "D1"), ("E1", "F1")):
        ws.merge_cells(f"{a}:{b}")
    for i, p in enumerate(PRODUCTS[3:11]):
        d = base - timedelta(days=rng.randint(0, 25))
        updated = d if i % 2 == 0 else d.strftime("%b %-d") if sys.platform != "win32" else d.strftime("%b ") + str(d.day)
        ws.append([
            f"PC {200 + i}",
            p.title() if i % 2 else p,
            f"${rng.uniform(5, 2500):,.2f}",
            "USD",
            rng.randint(0, 900),
            updated,
        ])
    wb.save(OUT / "pacific_stock.xlsx")


def lakeside(rng: random.Random, base: date) -> None:
    from fpdf import FPDF

    header = ["SKU", "Product", "Size / Color", "Price", "Qty", "Updated"]
    widths = [28, 62, 34, 24, 16, 26]
    sizes, colors = ["S", "M", "L", "XL"], ["navy", "Sand", "OLIVE", "white"]
    rows = []
    for i in range(30):
        p = PRODUCTS[i % len(PRODUCTS)]
        d = base - timedelta(days=rng.randint(0, 30))
        rows.append([f"LS-{300 + i}", p, f"{rng.choice(sizes)} / {rng.choice(colors)}",
                     f"${rng.uniform(3, 180):.2f}", str(rng.randint(0, 250)), d.isoformat()])

    pdf = FPDF(format="Letter")
    pdf.set_auto_page_break(False)
    pdf.set_font("Helvetica", size=9)

    def page(chunk):
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 10, "Lakeside Home Co. - Wholesale catalog (fictitious)", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "B", 9)
        for h, w in zip(header, widths):
            pdf.cell(w, 7, h, border=1)
        pdf.ln()
        pdf.set_font("Helvetica", size=9)
        for r in chunk:
            for v, w in zip(r, widths):
                pdf.cell(w, 6, v, border=1)
            pdf.ln()

    page(rows[:18])
    page(rows[18:])
    pdf.output(str(OUT / "lakeside_catalog.pdf"))


def maple(rng: random.Random, base: date) -> None:
    rows = [["Item #", "Name", "Variant", "Price (USD)", "Stock", "Last update"]]
    fmts = [lambda d: d.strftime("%m/%d/%y"), lambda d: d.isoformat(), lambda d: d.strftime("%B ") + str(d.day)]
    for i, p in enumerate(PRODUCTS[:10]):
        name, _, variant = p.rpartition(" ") if p[-1].isdigit() or p.endswith("l") else (p, "", "")
        d = base - timedelta(days=rng.randint(0, 25))
        rows.append([
            f"mp {400 + i}",
            name.upper() if i % 4 == 0 else name,
            variant,
            f"{rng.uniform(5, 1500):,.2f}",
            f" {rng.randint(0, 300)} ",
            fmts[i % 3](d),
        ])
    # Same SKU sent twice: an older row first, then the current one.
    old = list(rows[2])
    old[3] = "9.99"
    old[5] = (base - timedelta(days=60)).strftime("%m/%d/%y")
    rows.insert(2, old)
    with open(OUT / "maple_supplies.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def main(seed: int = 11, base: date = date(2026, 4, 20)) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    rng = _rng(seed)
    nordhaus(rng, base)
    pacific(rng, base)
    lakeside(rng, base)
    maple(rng, base)
    return OUT


if __name__ == "__main__":
    out = main()
    for p in sorted(out.iterdir()):
        print(f"  {p.relative_to(ROOT)}  {p.stat().st_size:,} bytes")
